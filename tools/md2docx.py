"""把 README.md 转成 Word（.docx）。

覆盖这份 README 用到的 Markdown 子集：标题、段落、行内 **粗体** / `代码` /
[链接](url)（含 #锚点 → Word 书签）、无序/有序列表、GitHub 表格、
围栏代码块、引用块、分隔线、本地图片（段落内或表格单元格里）。

用法：python md2docx.py <in.md> <out.docx>

依赖：python-docx（**不在 requirements.lock 里** —— 这是文档工具，不是服务依赖；
要跑先 `pip install python-docx`）。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"
MONO_FONT = "Consolas"

LINK_COLOR = RGBColor(0x11, 0x55, 0xCC)
CODE_COLOR = RGBColor(0xB4, 0x1E, 0x1E)
HEAD_COLOR = RGBColor(0x0B, 0x3D, 0x91)
GREY = RGBColor(0x77, 0x77, 0x77)

_bookmark_id = [100]
_slugs: dict[str, str] = {}

INLINE_RE = re.compile(
    r"(?P<img>!\[(?P<ialt>[^\]]*)\]\((?P<isrc>[^)]+)\))"
    r"|(?P<link>\[(?P<lt>[^\]]+)\]\((?P<lu>[^)]+)\))"
    r"|\*\*(?P<b>[^*]+)\*\*"
    r"|`(?P<c>[^`]+)`"
    r"|~~(?P<s>[^~]+)~~"
)


def slugify(text: str) -> str:
    t = text.strip().lower()
    t = re.sub(r"[`*_~\[\]()（）【】。，、：:；;！!？?“”\"'’‘—\-\./\\|]+", "", t)
    t = re.sub(r"\s+", "-", t)
    return t


def set_run_font(run, font: str = BODY_FONT, size: float | None = None,
                 bold: bool | None = None, color: RGBColor | None = None):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(paragraph, fill: str):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def add_bookmark(paragraph, name: str):
    _bookmark_id[0] += 1
    bid = str(_bookmark_id[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    paragraph._p.insert(0, start)
    paragraph._p.append(end)
    _slugs[name] = name


def add_hyperlink(paragraph, text: str, url: str, size: float):
    if url.startswith("#"):
        slug = slugify(url[1:])
        anchor = slug if slug in _slugs else None
        if anchor is None:
            run = paragraph.add_run(text)
            set_run_font(run, BODY_FONT, size, color=RGBColor(0x33, 0x33, 0x33))
            return
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), anchor)
        link.set(qn("w:history"), "1")
    else:
        r_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), r_id)
    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run_el.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run_el.append(t)
    link.append(run_el)
    paragraph._p.append(link)


def emit(paragraph, text: str, size: float = 10.5, bold: bool = False,
         color: RGBColor | None = None, base: Path | None = None,
         image_width: Inches | None = None):
    """把行内 Markdown 写进段落；粗体内部递归，保证「**…[链接](…)…**」也成链。"""
    text = re.sub(r"<br\s*/?>", "", text)
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos:m.start()])
            set_run_font(run, BODY_FONT, size, bold=bold, color=color)
        if m.group("img"):
            alt, src = m.group("ialt"), m.group("isrc")
            path = (base or Path(".")) / src
            if path.exists() and image_width is not None:
                run = paragraph.add_run()
                run.add_picture(str(path), width=image_width)
            else:
                run = paragraph.add_run(f"[图片：{alt}]")
                set_run_font(run, BODY_FONT, size - 1, color=GREY)
        elif m.group("link"):
            add_hyperlink(paragraph, m.group("lt"), m.group("lu"), size)
        elif m.group("b"):
            emit(paragraph, m.group("b"), size, bold=True, color=color,
                 base=base, image_width=image_width)
        elif m.group("c"):
            run = paragraph.add_run(m.group("c"))
            set_run_font(run, MONO_FONT, size - 0.5, color=CODE_COLOR)
        elif m.group("s"):
            run = paragraph.add_run(m.group("s"))
            set_run_font(run, BODY_FONT, size, color=color)
            run.font.strike = True
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, BODY_FONT, size, bold=bold, color=color)


# ---------------------------------------------------------------- 骨架
def new_document() -> Document:
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)
    return doc


def usable_width(doc):
    s = doc.sections[0]
    return s.page_width - s.left_margin - s.right_margin


def add_title(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    set_run_font(run, HEAD_FONT, 17, bold=True, color=RGBColor(0x1A, 0x1A, 0x1A))


def add_heading(doc, text: str, level: int):
    size = {2: 14, 3: 11.5, 4: 11}.get(level, 11)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 2 else 10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    emit(p, text, size=size)
    for run in p.runs:
        set_run_font(run, HEAD_FONT, size, bold=True, color=HEAD_COLOR)
    if level == 2:
        ppr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), "C9D6E8")
        borders.append(bottom)
        ppr.append(borders)
    add_bookmark(p, slugify(re.sub(r"^#+\s*", "", text)))


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "DDDDDD")
    borders.append(bottom)
    ppr.append(borders)


def add_quote(doc, lines: list[str]):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "9DB7D9")
    borders.append(left)
    ppr.append(borders)
    for idx, line in enumerate(lines):
        if idx:
            p.add_run().add_break()
        emit(p, line, size=10, color=RGBColor(0x44, 0x44, 0x44))


def add_code_block(doc, code: str, lang: str):
    # 代码块里不该出现 Markdown 的换行标签与 HTML 转义；Word 里直接显示成空白/原字符
    code = code.replace("<br/>", " ").replace("<br />", " ").replace("<br>", " ")
    code = code.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.25)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    shade(p, "F4F6F8")
    ppr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "E1E5EA")
        borders.append(el)
    ppr.append(borders)
    if lang:
        cap = p.add_run(f"[{lang}]\n")
        set_run_font(cap, MONO_FONT, 8, color=RGBColor(0x8A, 0x93, 0x9E))
    run = p.add_run(code)
    set_run_font(run, MONO_FONT, 9, color=RGBColor(0x24, 0x29, 0x2F))


def add_list_item(doc, text: str, level: int = 0, marker: str = "•"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.85 + 0.6 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(3)
    if marker and text[:1] not in "☐☑":
        text = f"{marker} {text}"
    emit(p, text, size=10.5)
    return p


def add_image(doc, path: Path, alt: str, width: Inches):
    if not path.exists():
        p = doc.add_paragraph()
        emit(p, f"[缺图：{alt}]", size=10, color=GREY)
        return
    doc.add_picture(str(path), width=width)
    pic = doc.paragraphs[-1]
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.space_before = Pt(4)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(alt)
    set_run_font(run, BODY_FONT, 8.5, color=GREY)


def add_figure_list(doc, rows: list[list[str]], base: Path, img_width: Inches):
    """截图表格 → 「小标题 + 居中大图」的竖排图组（Word 里比塞进表格干净得多）。

    按**列**配对：表头那行是说明，下一行是 `![alt](路径)`，所以第 c 列 =
    「说明 + 图」一组，而不是先把所有说明写一行、再堆所有图。
    """
    has_image_row = any("![" in cell for row in rows[1:] for cell in row)
    if has_image_row and len(rows) >= 2:
        cols = max(len(r) for r in rows)
        groups = []
        for c in range(cols):
            caption = rows[0][c].strip() if c < len(rows[0]) else ""
            imgs = []
            for row in rows[1:]:
                if c < len(row) and "![" in row[c]:
                    imgs.append(row[c].strip())
            groups.append((caption, imgs))
    else:
        groups = [(cell.strip(), [cell.strip()]) if "![" in cell else (cell.strip(), [])
                  for row in rows for cell in row]

    for caption, imgs in groups:
        if caption and "![" not in caption:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            emit(p, caption, size=10, bold=True)
        for blob in imgs:
            m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", blob)
            if not m:
                continue
            alt, src = m.group(1), m.group(2)
            path = base / src
            if path.exists():
                # 按「最多 3.4 英寸高」反推宽度，避免一张截图占满一页、留一大片空白
                try:
                    from PIL import Image as _Img
                    w_px, h_px = _Img.open(path).size
                    max_h = Inches(3.4)
                    w_in = min(img_width, int(max_h * w_px / h_px))
                except Exception:
                    w_in = img_width
                doc.add_picture(str(path), width=w_in)
                pic = doc.paragraphs[-1]
                pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.paragraph_format.space_before = Pt(2)
                pic.paragraph_format.space_after = Pt(10)
            else:
                p = doc.add_paragraph()
                emit(p, f"[缺图：{alt}]", size=9, color=GREY)


def add_table(doc, rows: list[list[str]], base: Path, cell_img_width: Inches):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_i, row in enumerate(rows):
        for c_i in range(cols):
            cell = table.cell(r_i, c_i)
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(1)
            par.paragraph_format.line_spacing = 1.15
            value = row[c_i] if c_i < len(row) else ""
            emit(par, value, size=9, bold=(r_i == 0), base=base,
                 image_width=cell_img_width)
            if r_i == 0:
                shade(par, "EDF2F9")
            for run in par.runs:
                if r_i == 0:
                    set_run_font(run, BODY_FONT, 9, bold=True)
        # 行不跨页；表头行在跨页后重复
        trpr = table.rows[r_i]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trpr.append(cant)
        if r_i == 0:
            hdr = OxmlElement("w:tblHeader")
            trpr.append(hdr)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- 解析
def parse_table_block(lines: list[str]) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line.startswith("|"):
            break
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            i += 1
            continue
        rows.append(cells)
        i += 1
    return rows, i


def convert(md_path: Path, out_path: Path):
    text = md_path.read_text(encoding="utf-8")
    lines = text.split("\n")
    base = md_path.parent
    doc = new_document()
    width = usable_width(doc)
    img_width = min(width, Inches(6.3))

    # 先扫一遍标题，建立锚点表（TOC 在正文之前，必须先有全书签）
    for line in lines:
        m = re.match(r"^(#{1,4})\s+(.*)$", line.strip())
        if m:
            _slugs[slugify(m.group(2).strip())] = slugify(m.group(2).strip())

    i = 0
    buf: list[str] = []
    last_item = [None]

    def flush():
        nonlocal buf
        item, last_item[0] = last_item[0], None
        if not buf:
            return
        joined = " ".join(x.strip() for x in buf).strip()
        buf = []
        if joined:
            p = doc.add_paragraph()
            emit(p, joined, size=10.5, base=base)

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()

        if stripped.startswith("```"):
            flush()
            lang = stripped[3:].strip()
            j = i + 1
            code = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j])
                j += 1
            add_code_block(doc, "\n".join(code), lang)
            i = j + 1
            continue

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            flush()
            add_rule(doc)
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            flush()
            level = len(m.group(1))
            title = m.group(2).strip()
            if level == 1:
                add_title(doc, title)
            else:
                add_heading(doc, title, level)
            i += 1
            continue

        if stripped.startswith("|"):
            flush()
            rows, consumed = parse_table_block(lines[i:])
            has_image = any("![" in cell for row in rows for cell in row)
            if has_image:
                add_figure_list(doc, rows, base, min(width, Inches(5.9)))
            else:
                cell_w = Inches(3.0) if max(len(r) for r in rows) <= 2 else Inches(2.0)
                add_table(doc, rows, base, cell_w)
            i += consumed
            continue

        m_img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
        if m_img:
            flush()
            add_image(doc, base / m_img.group(2), m_img.group(1), img_width)
            i += 1
            continue

        if stripped.startswith(">"):
            flush()
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            add_quote(doc, [q for q in quote if q != ""] or [""])
            continue

        m_ul = re.match(r"^[-*]\s+(.*)$", stripped)
        m_ol = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_ul or m_ol:
            flush()
            indent = len(raw) - len(raw.lstrip())
            if m_ul:
                body, marker = m_ul.group(1), "•"
            else:
                body, marker = m_ol.group(2), f"{m_ol.group(1)}."
            body = body.replace("[ ]", "☐ ").replace("[x]", "☑ ")
            # 目录行（整行就是一个 #锚点 链接）不加圆点，读起来像目录而不是清单
            if re.fullmatch(r"\[[^\]]+\]\(#[^)]+\)", body.strip()):
                marker = ""
            last_item[0] = add_list_item(
                doc, body, level=0 if indent < 2 else 1, marker=marker)
            i += 1
            continue

        if stripped == "":
            flush()
            i += 1
            continue

        # 列表项的续行（Markdown 里用缩进续写同一项）：接回上一个列表段落
        if (last_item[0] is not None and stripped
                and raw[:1] in (" ", "\t")
                and not re.match(r"^([-*]\s|\d+\.\s|#|\||>|```)", stripped)):
            last_item[0].add_run(" ")
            emit(last_item[0], stripped, size=10.5, base=base)
            i += 1
            continue

        buf.append(stripped)
        i += 1

    flush()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


if __name__ == "__main__":
    print(convert(Path(sys.argv[1]), Path(sys.argv[2])))
